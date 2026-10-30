#!/usr/bin/env python3
"""Lemma 問題管理表（.xlsx）→ アプリ用データ（data/*.json）

使い方：
    pip3 install openpyxl
    python3 tools/import_xlsx.py Lemma問題管理表.xlsx

出力：
    data/index.json        全問の一覧（ID・科目・分野・難易度・目安時間・タグ・テーマ）
    data/f-<分野>.json      分野ごとの本文（問題文・答え・解説・核となるアイデア）

「参考元」「作成方法」「メモ」はアプリに出さないため、出力に含めない。
問題文などの LaTeX は、文章部分（enumerate・itemize・tabular・\\textbf など）を HTML に直し、
数式部分（$...$、\\[...\\]）はそのまま残す（アプリ側で KaTeX が描画する）。
"""
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit('openpyxl が必要です： pip3 install openpyxl')

SUBJECT_KEY = {'数I': 'I', '数A': 'A', '数II': 'II', '数B': 'B', '数III': 'III', '数C': 'C', '融合': 'X'}
COLS = ['id', 'date', 'subject', 'field', 'theme', 'level', 'minutes', 'q', 'answer', 'solution',
        'idea', 'source', 'method', 'v1', 'v2', 'v3', 'verified', 'tags', 'memo']

WARN = []


# ---------------------------------------------------------------- LaTeX → HTML
def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


DISPLAY_ENVS = {'align*': 'aligned', 'align': 'aligned', 'gather*': 'gathered', 'gather': 'gathered',
                'equation*': None, 'equation': None, 'eqnarray*': 'aligned'}


def read_group(s, i):
    """s[i] が '{' のとき、対応する '}' までの中身と次の位置を返す"""
    if i >= len(s) or s[i] != '{':
        return None, i
    depth, j = 0, i
    while j < len(s):
        c = s[j]
        if c == '\\':
            j += 2
            continue
        if c == '{':
            depth += 1
        elif c == '}':
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    return s[i + 1:], len(s)


def tokenize(s):
    """文章を [('math', raw, display), ('text', str), ('begin', env, arg), ('end', env), ('item', label), ('br',), ('par',)] に分ける"""
    toks, buf, i, n = [], [], 0, len(s)

    def flush():
        if buf:
            toks.append(('text', ''.join(buf)))
            buf.clear()

    while i < n:
        if s.startswith('\\\\', i):
            flush()
            i += 2
            m = re.match(r'\s*\[[^\]]*\]', s[i:])  # \\[2pt] などの余白指定
            if m:
                i += m.end()
            toks.append(('br',))
            continue
        if s.startswith('\\$', i):
            buf.append('$')
            i += 2
            continue
        if s[i] == '$':
            j = i + 1
            while j < n and not (s[j] == '$' and s[j - 1] != '\\'):
                m2 = re.match(r'\\(?:text|mbox|textrm|textbf|hbox)\s*(?=\{)', s[j:])
                if m2:  # \text{...} の中の $ は数式の終わりではない
                    _, j = read_group(s, j + m2.end())
                    continue
                if s[j] == '\\':
                    j += 2
                    continue
                j += 1
            flush()
            toks.append(('math', s[i + 1:j], False))
            i = j + 1
            continue
        if s.startswith('\\[', i) or s.startswith('\\(', i):
            close = '\\]' if s[i + 1] == '[' else '\\)'
            j = s.find(close, i + 2)
            if j < 0:
                j = n
            flush()
            toks.append(('math', s[i + 2:j], close == '\\]'))
            i = j + 2
            continue
        m = re.match(r'\\begin\{([a-zA-Z]+\*?)\}', s[i:])
        if m:
            env = m.group(1)
            if env in DISPLAY_ENVS:
                endtag = '\\end{%s}' % env
                j = s.find(endtag, i + m.end())
                if j < 0:
                    j = n
                body = s[i + m.end():j]
                inner = DISPLAY_ENVS[env]
                flush()
                toks.append(('math', ('\\begin{%s}%s\\end{%s}' % (inner, body, inner)) if inner else body, True))
                i = j + len(endtag)
                continue
            flush()
            i += m.end()
            arg = None
            if env == 'tabular':
                arg, i = read_group(s, i)
            toks.append(('begin', env, arg))
            continue
        m = re.match(r'\\end\{([a-zA-Z]+\*?)\}', s[i:])
        if m:
            flush()
            toks.append(('end', m.group(1)))
            i += m.end()
            continue
        m = re.match(r'\\item\b\s*', s[i:])
        if m:
            flush()
            i += m.end()
            label = None
            if i < n and s[i] == '[':
                j = s.find(']', i)
                label = s[i + 1:j]
                i = j + 1
            toks.append(('item', label))
            continue
        m = re.match(r'\\setcounter\{enumi\}\{(\d+)\}', s[i:])
        if m:
            flush()
            toks.append(('counter', int(m.group(1))))
            i += m.end()
            continue
        if s.startswith('\n', i):
            j = i
            while j < n and s[j] in ' \t\n':
                j += 1
            flush()
            toks.append(('par',) if s[i:j].count('\n') >= 2 else ('nl',))
            i = j
            continue
        buf.append(s[i])
        i += 1
    flush()
    return toks


TEXT_CMDS = {
    'quad': '　', 'qquad': '　　', ',': ' ', ' ': ' ', ';': ' ', ':': ' ', 'hfill': '<span class="hfill"></span>',
    'noindent': '', 'par': '<br>', 'medskip': '', 'smallskip': '', 'bigskip': '', 'newline': '<br>',
    'ldots': '…', 'dots': '…', 'cdots': '⋯', '%': '%', '&': '&amp;', '#': '#', '_': '_', '{': '{', '}': '}',
    'S': '§', 'textbackslash': '\\',
}
WRAP_CMDS = {'textbf': 'b', 'bf': 'b', 'emph': 'em', 'textit': 'em', 'underline': 'u', 'textrm': None, 'mbox': None,
             'text': None, 'textgt': 'b', 'textsf': None}


def text_inline(t, ctx):
    """数式でない文章片を HTML に（\\textbf などの命令を処理）"""
    out, i, n = [], 0, len(t)
    while i < n:
        c = t[i]
        if c == '\\':
            m = re.match(r'\\([a-zA-Z]+|.)', t[i:])
            name = m.group(1)
            j = i + m.end()
            if name in WRAP_CMDS:
                while j < n and t[j] == ' ':
                    j += 1
                body, j = read_group(t, j)
                tag = WRAP_CMDS[name]
                inner = text_inline(body or '', ctx)
                out.append(f'<{tag}>{inner}</{tag}>' if tag else inner)
            elif name == 'fbox':
                body, j = read_group(t, j)
                body = body or ''
                if re.fullmatch(r'\s*\\phantom\{[^}]*\}\s*', body):
                    out.append('<span class="fbox blank"></span>')
                else:
                    out.append(f'<span class="fbox">{text_inline(body, ctx)}</span>')
            elif name == 'parbox':
                _, j = read_group(t, j)       # 幅
                body, j = read_group(t, j)
                out.append(f'<span class="parbox">{render(body or "", ctx)}</span>')
            elif name == 'phantom':
                body, j = read_group(t, j)
                out.append(f'<span class="phantom">{text_inline(body or "", ctx)}</span>')
            elif name in ('vspace', 'hspace', 'vspace*', 'hspace*'):
                if j < n and t[j] == '*':
                    j += 1
                _, j = read_group(t, j)
            elif name in TEXT_CMDS:
                out.append(TEXT_CMDS[name])
            elif name == 'multicolumn':
                _, j = read_group(t, j)
                _, j = read_group(t, j)
                body, j = read_group(t, j)
                out.append(text_inline(body or '', ctx))
            else:
                ctx['unknown'].add(name)
                out.append(esc(name))
            i = j
            continue
        if c in '{}':
            i += 1
            continue
        if c == '~':
            out.append('&nbsp;')
            i += 1
            continue
        out.append(esc(c))
        i += 1
    return ''.join(out)


def math_html(raw, display):
    raw = re.sub(r'\s*\n\s*', ' ', raw)  # 数式中の改行は空白に（<br> にすると数式が切れる）
    raw = raw.strip() if display else raw
    body = esc(raw)
    return f'<div class="dm">\\[{body}\\]</div>' if display else f'${body}$'


def build(toks, ctx, stop=None):
    """トークン列を HTML に。stop の \\end が来たら戻る"""
    html = []
    while ctx['pos'] < len(toks):
        tk = toks[ctx['pos']]
        ctx['pos'] += 1
        kind = tk[0]
        if kind == 'text':
            html.append(text_inline(tk[1], ctx))
        elif kind == 'math':
            html.append(math_html(tk[1], tk[2]))
        elif kind == 'br':
            html.append('<br>')
        elif kind == 'nl':
            html.append('\n')
        elif kind == 'par':
            html.append('\n\n')
        elif kind == 'end':
            if tk[1] == stop:
                return ''.join(html)
            WARN.append(f'{ctx["id"]}: 対応しない \\end{{{tk[1]}}}')
        elif kind == 'begin':
            env, arg = tk[1], tk[2]
            if env in ('enumerate', 'itemize'):
                html.append(build_list(toks, ctx, env))
            elif env == 'tabular':
                html.append(build_tabular(toks, ctx))
            elif env in ('center', 'flushleft', 'flushright', 'quote', 'quotation', 'minipage'):
                inner = build(toks, ctx, env)
                cls = {'center': 'center', 'quote': 'quote', 'quotation': 'quote'}.get(env, '')
                html.append(f'<div class="{cls}">{tidy(inner)}</div>')
            else:
                WARN.append(f'{ctx["id"]}: 未対応の環境 {env}（中身はそのまま表示）')
                html.append(build(toks, ctx, env))
        elif kind == 'item':
            WARN.append(f'{ctx["id"]}: リスト外の \\item')
        elif kind == 'counter':
            pass
    return ''.join(html)


def build_list(toks, ctx, env):
    tag = 'ol' if env == 'enumerate' else 'ul'
    items, cur, nextval, num = [], None, None, 0
    while ctx['pos'] < len(toks):
        tk = toks[ctx['pos']]
        if tk[0] == 'end' and tk[1] == env:
            ctx['pos'] += 1
            break
        if tk[0] == 'item':
            ctx['pos'] += 1
            num = nextval if nextval else num + 1
            nextval = None
            cur = {'label': tk[1], 'num': num, 'parts': []}
            items.append(cur)
            continue
        if tk[0] == 'counter':
            ctx['pos'] += 1
            nextval = tk[1] + 1
            continue
        # 項目の中身（入れ子リストも含め1要素ずつ処理）
        sub = {'pos': 0, 'id': ctx['id'], 'unknown': ctx['unknown']}
        if tk[0] == 'begin':
            ctx['pos'] += 1
            seg = build([tk] + collect_until_end(toks, ctx, tk[1]), sub)
        else:
            ctx['pos'] += 1
            seg = build([tk], sub)
        if cur is None:
            if seg.strip():
                WARN.append(f'{ctx["id"]}: \\item の前に文字があります')
            continue
        cur['parts'].append(seg)
    lis = []
    for it in items:
        body = tidy(''.join(it['parts']))
        if it['label'] is not None:
            lab = text_inline(it['label'], ctx)
            lis.append(f'<li class="lab"><span class="lb">{lab}</span>{body}</li>')
        elif tag == 'ol':
            lis.append(f'<li value="{it["num"]}"><span class="lb">({it["num"]})</span>{body}</li>')
        else:
            lis.append(f'<li>{body}</li>')
    return f'<{tag}>{"".join(lis)}</{tag}>'


def collect_until_end(toks, ctx, env):
    """toks[ctx.pos:] から、対応する \\end{env} までを取り出す（入れ子対応）"""
    depth, out = 1, []
    while ctx['pos'] < len(toks):
        tk = toks[ctx['pos']]
        ctx['pos'] += 1
        out.append(tk)
        if tk[0] == 'begin' and tk[1] == env:
            depth += 1
        elif tk[0] == 'end' and tk[1] == env:
            depth -= 1
            if depth == 0:
                break
    return out


def build_tabular(toks, ctx):
    rows, cells, cur = [], [], []
    def end_cell():
        cells.append(''.join(cur).strip()); cur.clear()
    while ctx['pos'] < len(toks):
        tk = toks[ctx['pos']]
        ctx['pos'] += 1
        if tk[0] == 'end' and tk[1] == 'tabular':
            break
        if tk[0] == 'br':
            end_cell(); rows.append(cells[:]); cells.clear()
            continue
        if tk[0] == 'math':
            cur.append(math_html(tk[1], False))
            continue
        if tk[0] == 'text':
            parts = re.sub(r'\\hline|\\cline\{[^}]*\}', '', tk[1]).split('&')
            for k, part in enumerate(parts):
                if k:
                    end_cell()
                cur.append(text_inline(part, ctx))
            continue
    if ''.join(cur).strip() or cells:
        end_cell(); rows.append(cells[:])
    rows = [r for r in rows if any(c for c in r)]
    body = ''.join('<tr>' + ''.join(f'<td>{c}</td>' for c in r) + '</tr>' for r in rows)
    return f'<div class="tbl"><table>{body}</table></div>'


BLOCK = r'(?:ol|ul|li|div|table|tr|td)'


def tidy(h):
    """改行を <br> にし、ブロック要素の前後の余分な改行を除く"""
    h = h.strip()
    h = re.sub(r'\n{2,}', '<br class="p">', h)
    h = h.replace('\n', '<br>')
    h = re.sub(r'(<br[^>]*>\s*)+(?=</?' + BLOCK + r'\b)', '', h)
    h = re.sub(r'(</?' + BLOCK + r'\b[^>]*>)\s*(<br[^>]*>\s*)+', r'\1', h)
    h = re.sub(r'^(<br[^>]*>)+|(<br[^>]*>)+$', '', h)
    return h


def render(src, ctx):
    toks = tokenize(str(src).replace('\r\n', '\n').replace('\r', '\n'))
    sub = {'pos': 0, 'id': ctx['id'], 'unknown': ctx['unknown']}
    return tidy(build(toks, sub))


# ---------------------------------------------------------------- 読み込み
def field_key(subject, field, keys):
    k = keys.get((subject, field))
    if not k:
        k = f'{SUBJECT_KEY.get(subject, "X")}-{len(keys) + 1:02d}'
        keys[(subject, field)] = k
    return k


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).resolve().parent.parent / 'data'
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb['問題一覧']
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    # 分野の並び順は「分野一覧」シートに従う
    order = []
    if '分野一覧' in wb.sheetnames:
        for r in wb['分野一覧'].iter_rows(min_row=2, values_only=True):
            if r[0] and r[1] and r[0] in SUBJECT_KEY:
                order.append((r[0], r[1]))
    keys = {}
    for sf in order:
        field_key(*sf, keys)

    index, bodies, seen, unknown = [], {}, set(), set()
    for r in rows:
        if not r or not r[0]:
            continue
        rec = dict(zip(COLS, list(r) + [None] * (len(COLS) - len(r))))
        pid = str(rec['id']).strip()
        if pid in seen:
            WARN.append(f'{pid}: ID が重複しています（2件目以降は無視）')
            continue
        seen.add(pid)
        subject, field = str(rec['subject'] or '').strip(), str(rec['field'] or '').strip()
        if subject not in SUBJECT_KEY:
            WARN.append(f'{pid}: 科目「{subject}」が不明なので飛ばします')
            continue
        if not rec['q'] or not rec['answer']:
            WARN.append(f'{pid}: 問題文か答えが空なので飛ばします')
            continue
        fk = field_key(subject, field, keys)
        ctx = {'id': pid, 'unknown': unknown}
        try:
            lv = int(rec['level'])
        except (TypeError, ValueError):
            lv = 3
            WARN.append(f'{pid}: 難易度が数でないので 3 にしました')
        lv = min(5, max(1, lv))
        try:
            mins = int(rec['minutes']) if rec['minutes'] else None
        except (TypeError, ValueError):
            mins = None
        tags = [t.strip() for t in re.split(r'[,、，]', str(rec['tags'] or '')) if t.strip()]
        entry = {'id': pid, 'f': fk, 'lv': lv}
        if mins:
            entry['t'] = mins
        if rec['theme']:
            entry['th'] = render(rec['theme'], ctx)
        if tags:
            entry['tg'] = tags
        if rec['verified'] == '✓' or (rec['v1'] == '済' and rec['v2'] == '済' and rec['v3'] == '済'):
            entry['ok'] = 1
        index.append(entry)
        bodies.setdefault(fk, {})[pid] = {
            'q': render(rec['q'], ctx),
            'a': render(rec['answer'], ctx),
            's': render(rec['solution'] or '', ctx),
            'k': render(rec['idea'] or '', ctx),
        }

    out.mkdir(parents=True, exist_ok=True)
    for old in out.glob('f-*.json'):
        old.unlink()
    fields = [{'key': k, 'subject': s, 'name': f} for (s, f), k in keys.items()
              if any(e['f'] == k for e in index)]
    meta = {'generated': datetime.now().isoformat(timespec='seconds'), 'source': src.name,
            'count': len(index), 'fields': fields, 'problems': index}
    (out / 'index.json').write_text(json.dumps(meta, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    for fk, b in bodies.items():
        (out / f'f-{fk}.json').write_text(json.dumps(b, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')

    print(f'{len(index)}問を書き出しました（{len(fields)}分野） → {out}')
    if unknown:
        print('文章中の未対応命令（名前だけ表示）:', ', '.join(sorted(unknown)))
    if WARN:
        print(f'注意 {len(WARN)}件:')
        for w in WARN[:50]:
            print('  ', w)


if __name__ == '__main__':
    main()
